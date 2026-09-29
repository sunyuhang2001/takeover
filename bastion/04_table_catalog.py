# -*- coding: utf-8 -*-
# 04_table_catalog.py : 所有表格(xlsx/csv)只读表头，按列名分组，找出轨迹表（只读）
# 依赖 01_scan_files 生成的 file_index.csv；需要 pandas、openpyxl
import os, sys, re, csv, hashlib, time
from collections import defaultdict
import pandas as pd

NAME, VERSION = "04_table_catalog", "v1"
INDEX = r"D:\takeover_audit\01_scan\file_index.csv"
OUT_DIR = r"D:\takeover_audit\04_tables"
MAX_XLSX_MB = 300   # 超过这个大小的 xlsx 跳过（记录下来），避免卡死


def selfcheck():
    try:
        t = open(__file__, "rb").read().decode("utf-8", "replace")
    except Exception:
        return "n/a"
    lines = [l.rstrip() for l in t.replace("\r\n", "\n").replace("\r", "\n").split("\n")]
    while lines and lines[-1] == "":
        lines.pop()
    return hashlib.md5("\n".join(lines).encode("utf-8")).hexdigest()[:8]


L = []
def out(s=""):
    print(s); L.append(s)


def read_csv_head(path):
    for enc in ("utf-8-sig", "gbk"):
        try:
            with open(path, "r", encoding=enc, newline="") as f:
                rows = []
                for i, r in enumerate(csv.reader(f)):
                    rows.append(r)
                    if i >= 2:
                        break
            return [("", rows[0] if rows else [], rows[1] if len(rows) > 1 else [], None, enc)]
        except UnicodeDecodeError:
            continue
    raise ValueError("encoding")


def read_xlsx_head(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    res = []
    for ws in wb.worksheets:
        it = ws.iter_rows(max_row=2, values_only=True)
        hdr = next(it, ()) or ()
        ex = next(it, ()) or ()
        res.append((ws.title, ["" if h is None else str(h) for h in hdr],
                    ["" if x is None else str(x) for x in ex], ws.max_row, ""))
    wb.close()
    return res


def flags(cols):
    low = [c.strip().lower() for c in cols]
    j = " ".join(low)
    has = lambda *ks: any(any(k in c for k in ks) for c in low)
    return {
        "vin": has("vin", "车架"),
        "time": has("time", "时间", "timestamp"),
        "lat": has("lat", "纬度"), "lon": has("lon", "lng", "经度"),
        "drive_mode": has("drive_mode", "驾驶模式"),
        "desc": has("描述", "原因", "desc"),
        "ptc": has("ptcid", "participant", "目标id", "obj_id", "track_id"),
    }


def main():
    try:
        sys.stdout.reconfigure(errors="replace")
    except Exception:
        pass
    os.makedirs(OUT_DIR, exist_ok=True)
    out("=== %s %s  check=%s  %s ===" % (NAME, VERSION, selfcheck(), time.strftime("%Y-%m-%d %H:%M")))
    fi = pd.read_csv(INDEX, dtype=str, keep_default_na=False)
    fi = fi[(fi.kind == "table") & ~fi.dir.str.lower().str.contains("takeover_audit") & ~fi.name.str.startswith("~$")]
    fi["size_bytes"] = pd.to_numeric(fi.size_bytes, errors="coerce").fillna(0)
    recs, errs, t0 = [], [], time.time()
    for i, (d, n, ext, sz, top) in enumerate(zip(fi.dir, fi.name, fi.ext, fi.size_bytes, fi["top"])):
        path = d + "\\" + n
        try:
            if ext == ".csv":
                sheets = read_csv_head(path)
            elif ext in (".xlsx", ".xlsm"):
                if sz > MAX_XLSX_MB * 2**20:
                    errs.append((path, "skip_big_xlsx")); continue
                sheets = read_xlsx_head(path)
            else:
                errs.append((path, "skip_" + ext)); continue
        except Exception as e:
            errs.append((path, repr(e)[:200])); continue
        for sn, hdr, ex, nrow, enc in sheets:
            cols = [h.strip() for h in hdr if h is not None]
            f = flags(cols)
            recs.append(dict(path=path, top=top, sheet=sn, size_mb=sz / 2**20, nrow=nrow, ncol=len(cols),
                             sig="|".join(sorted(c.lower() for c in cols if c))[:1000],
                             cols=",".join(cols)[:1500], example=",".join(ex)[:800], **f))
        if (i + 1) % 500 == 0:
            print("  ... %d/%d  %.0fs" % (i + 1, len(fi), time.time() - t0))
    df = pd.DataFrame(recs)
    df["traj"] = df.vin & df.time & df.lat & df.lon
    df["evlist"] = df.vin & df.time & df.desc
    df.to_csv(os.path.join(OUT_DIR, "table_catalog.csv"), index=False, encoding="utf-8-sig")
    pd.DataFrame(errs, columns=["path", "err"]).to_csv(os.path.join(OUT_DIR, "errors.csv"), index=False, encoding="utf-8-sig")

    out("表格文件=%d  sheet=%d  读失败/跳过=%d  用时=%.0fs" % (len(fi), len(df), len(errs), time.time() - t0))
    er = pd.Series([e for _, e in errs]).str[:20].value_counts()
    out("跳过原因: " + "  ".join("%s:%d" % (k, v) for k, v in er.head(5).items()))
    out("轨迹候选(vin+时间+经纬度) sheet=%d  文件=%d  %.1fGB ;  事件清单候选(vin+时间+描述)=%d ;  含ptcId类=%d" % (
        df.traj.sum(), df[df.traj].path.nunique(), df[df.traj].drop_duplicates("path").size_mb.sum() / 1024,
        df.evlist.sum(), df.ptc.sum()))

    # 按列名签名分组，只列轨迹候选
    sig = df[df.traj].groupby("sig").agg(files=("path", "nunique"), gb=("size_mb", lambda s: s.sum() / 1024),
                                          rows=("nrow", lambda s: s.dropna().sum()),
                                          cols=("cols", "first"), eg=("path", "first"),
                                          dm=("drive_mode", "first")).sort_values("gb", ascending=False)
    out("")
    out("轨迹候选按表结构分组 (前15组): 文件数 / GB / xlsx行数合计 / 有drive_mode")
    for k, (_, r) in enumerate(sig.head(15).iterrows(), 1):
        out("[%d] files=%d  %.2fGB  rows=%s  dm=%s  eg=%s" % (k, r.files, r.gb, int(r.rows) if r.rows else "-", r.dm, r.eg[-70:]))
        out("     cols: " + r.cols[:230])
    sig.to_csv(os.path.join(OUT_DIR, "traj_signatures.csv"), encoding="utf-8-sig")
    out("")
    out("轨迹候选所在一级目录: " + "  ".join("%s:%d" % (k.split("\\")[-1], v) for k, v in df[df.traj].drop_duplicates("path").top.value_counts().head(10).items()))
    out("=== END %s %s  -> %s ===" % (NAME, VERSION, OUT_DIR))
    open(os.path.join(OUT_DIR, "summary.txt"), "w", encoding="utf-8").write("\n".join(L))


if __name__ == "__main__":
    main()
